# -*- coding: utf-8 -*-
"""Dump the two sample Excel files (openpyxl, 1-based) used by the legacy towerdesign tool.

Mapping to legacy xlrd (0-based):  openpyxl(r+1, c+1) == xlrd.cell_value(r, c)
"""
import openpyxl

GEO = r"D:\work\Ontology\GoldenWind\design\towerdesign-main\TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
LAY = r"D:\work\Ontology\GoldenWind\design\towerdesign-main\项目布局表.xlsx"


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).replace("\n", "\\n")
    return s[:24] + "…" if len(s) > 24 else s


def dump(book, name, max_rows=45, max_cols=20):
    if name not in book.sheetnames:
        print(f"  [MISS] sheet '{name}'")
        return
    sh = book[name]
    print(f"  --- sheet '{name}'  {sh.max_row} rows x {sh.max_column} cols ---")
    for r in range(1, min(sh.max_row, max_rows) + 1):
        vals = [fmt(sh.cell(r, c).value) for c in range(1, min(sh.max_column, max_cols) + 1)]
        print(f"   r{r:>3}: " + " | ".join(vals))


def main():
    print("=" * 110)
    print("TowerGeoInput:", GEO.split("\\")[-1])
    print("=" * 110)
    gb = openpyxl.load_workbook(GEO, data_only=True)
    print("sheets:", gb.sheetnames)
    for nm in ["Description", "TowerGeo", "Flange", "Door"]:
        dump(gb, nm, max_rows=50, max_cols=18)

    print()
    print("=" * 110)
    print("项目布局表:", LAY.split("\\")[-1])
    print("=" * 110)
    lb = openpyxl.load_workbook(LAY, data_only=True)
    print("sheets:", lb.sheetnames)
    for nm in lb.sheetnames:
        if nm in ("中间段", "下段", "顶段"):
            dump(lb, nm, max_rows=40, max_cols=12)


if __name__ == "__main__":
    main()
