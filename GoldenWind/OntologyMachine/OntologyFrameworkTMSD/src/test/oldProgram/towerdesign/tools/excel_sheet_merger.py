import zipfile
import os
import pandas as pd
from openpyxl import load_workbook
import shutil


def merge_excel_sheets(zip_file_path, source_sheet_name, target_sheet_name, output_dir):
    temp_dir = 'temp_excel_files'
    os.makedirs(temp_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)  # 确保输出目录存在

    try:
        with zipfile.ZipFile(zip_file_path, 'r') as zip_ref:
            zip_ref.extractall(temp_dir)

        for file_name in os.listdir(temp_dir):
            if file_name.endswith('.xlsx'):
                file_path = os.path.join(temp_dir, file_name)
                output_file_path = os.path.join(output_dir, file_name)

                # 加载整个工作簿
                wb = load_workbook(file_path)
                sheet_names = wb.sheetnames

                # 尝试读取源sheet
                if source_sheet_name in sheet_names:
                    df_source = pd.read_excel(file_path, sheet_name=source_sheet_name)

                    # 如果目标sheet存在则移除
                    if target_sheet_name in sheet_names:
                        del wb[target_sheet_name]

                    # 创建新的目标sheet并写入数据
                    ws_target = wb.create_sheet(title=target_sheet_name)

                    # 确保列名不为 'Unnamed: X'
                    column_titles = df_source.columns.tolist()
                    cleaned_titles = [title if 'Unnamed' not in str(title) else f"Column{idx + 1}" for idx, title in
                                      enumerate(column_titles)]
                    ws_target.append(cleaned_titles)

                    # 写入数据到目标sheet
                    for row in df_source.itertuples(index=False):
                        ws_target.append(row)

                    # 将目标sheet移动到第一个位置
                    other_sheets = [ws for ws in wb.worksheets if ws.title != target_sheet_name]
                    wb._sheets = [ws_target] + other_sheets

                    # 保存修改后的工作簿
                    wb.save(file_path)

                    # 将处理后的文件移动到输出目录
                    shutil.move(file_path, output_file_path)
                    print(f"Processed file saved to: {output_file_path}")
                else:
                    print(f"Warning: {source_sheet_name} not found in {file_name}. Skipping this file.")
    finally:
        # 清理临时文件夹
        for file_name in os.listdir(temp_dir):
            file_path = os.path.join(temp_dir, file_name)
            if os.path.isfile(file_path):
                os.remove(file_path)
        os.rmdir(temp_dir)
        print("所有Excel文件处理完成！")

    # 获取zip_file所在的路径和文件名
    zip_file_dir = os.path.dirname(zip_file_path)
    zip_file_name = os.path.splitext(os.path.basename(zip_file_path))[0]
    output_zip_path = os.path.join(zip_file_dir, f"{zip_file_name}_processed_files.zip")

    # 打包输出目录为zip文件
    with zipfile.ZipFile(output_zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for root, _, files in os.walk(output_dir):
            for file in files:
                file_path = os.path.join(root, file)
                arcname = os.path.relpath(file_path, output_dir)
                zipf.write(file_path, arcname)
    print(f"所有文件已打包到: {output_zip_path}")



# 示例使用
if __name__ == "__main__":
    zip_file = r'D:\塔架计算平台\50.覆盖\Markov15_GWH221P11100H120_D1_GWBDA02_V8_2r8c_V26k24_HDHM_M4_tower_section.zip'  # 替换为你的zip文件路径
    source_sheet = 'My'  # 源sheet名称
    target_sheet = 'Mx'  # 目标sheet名称
    output_directory = r'D:\塔架计算平台\50.覆盖\output'  # 替换为你想要保存输出文件的目录

    merge_excel_sheets(zip_file, source_sheet, target_sheet, output_directory)
