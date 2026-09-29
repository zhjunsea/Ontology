from steeltowerdesign.tools.read_allMydata import all_flangeRange_markov_dict, single_flange_markov_dict
from steeltowerdesign.Flange.topFlangeFatigue import markov_xlsx2dict
from steeltowerdesign.tools.towerGeo2dict import tower_geo_2_dict
import openpyxl
import os
import numpy as np
import pandas as pd


def single_loadcase_damage(loc_flange_markovMatrix,load_name, load_spectrum, stress_spectrum, delta_sigmaD, Nd):
    #载荷
    my_mean_list = np.array(loc_flange_markovMatrix[load_name]["meanValues"]) #均值 mean
    delta_my_list = np.array(loc_flange_markovMatrix[load_name]["rangeValues"]) #差值 range
    # 将list转化为np.array
    my_mean_array = np.array(my_mean_list)
    delta_my_array = np.array(delta_my_list)
    #计算最大值与最小值
    my_mean_64x64T = my_mean_array * np.ones((my_mean_array.shape[0], my_mean_array.shape[0]))
    my_mean_64x64 = my_mean_64x64T.T  # 均值mean
    delta_my_64x64 = delta_my_array * np.ones((delta_my_array.shape[0], delta_my_array.shape[0]))  # Range
    my_max_64x64 = my_mean_64x64 + delta_my_64x64 / 2.0
    my_min_64x64 = my_mean_64x64 - delta_my_64x64 / 2.0

    # 应用到两个矩阵
    max_stress_64x64 = vectorized_interpolation(my_max_64x64, load_spectrum, stress_spectrum)
    min_stress_64x64 = vectorized_interpolation(my_min_64x64, load_spectrum, stress_spectrum)

    # 计算两个矩阵的差值
    stress_range_64x64 = abs(max_stress_64x64 - min_stress_64x64)

    # 计算N_lifeTime_i 64x64
    N_lifeTime_i64x64 = N_lifeTime(delta_sigmaD, stress_range_64x64, Nd)

    # 获得循环次数64x64
    number_i64x64 = np.array(loc_flange_markovMatrix[load_name]["metric"])
    damage = sum(sum(number_i64x64 / N_lifeTime_i64x64))
    #print(damage)
    return damage

def N_lifeTime(delta_sigmaD, delta_sigma_i, Nd):
    m = np.ones_like(delta_sigma_i) * 5
    try:
        m[delta_sigma_i >= delta_sigmaD] = 3
        life_time = np.ones_like(delta_sigma_i) * float('inf')
        life_time[delta_sigma_i != 0] = (Nd * (delta_sigmaD / delta_sigma_i) ** m)[delta_sigma_i != 0]
        return life_time
    except TypeError:
        if delta_sigma_i >= delta_sigmaD:
            m = 3
        if delta_sigma_i == 0:
            return float('inf')
        else:
            return Nd * (delta_sigmaD / delta_sigma_i) ** m
    except Exception as e:
        raise e

# 更高效的方式：使用numpy的interp函数
def vectorized_interpolation(matrix_to_interpolate, x_known, y_known):
    """
    使用numpy的向量化插值
    """
    # 展平矩阵进行插值
    flat_matrix = matrix_to_interpolate.flatten()
    interpolated_flat = np.interp(flat_matrix, x_known, y_known)

    # 重塑回原始形状
    interpolated_matrix = interpolated_flat.reshape(matrix_to_interpolate.shape)

    return interpolated_matrix

def calculate_bolt_fatigue_damage(towerGeoExcelName, markov, boltStressExcel, outputfile, elevation, mmboltType):
    #markovDir若是解压后的文件夹，则markovDir+"\\" + "markov"
    if os.path.isdir(markov):
        flange_worksheet = openpyxl.load_workbook(towerGeoExcelName, data_only=True)['Flange']
        flange_loc_list = []
        for i in range(3, flange_worksheet.max_row + 1):
            if flange_worksheet.cell(i, 1).value:
                flange_loc_list += [flange_worksheet.cell(i, 1).value]
        markovMatrix_dict = all_flangeRange_markov_dict(markov, flange_loc_list)
        loc_flange_markov_dict = single_flange_markov_dict(markovMatrix_dict, elevation)
    else:    #markovDir若是载荷excel
        loc_flange_markov_dict = markov_xlsx2dict(markov)
    #获取螺栓的规格
    if towerGeoExcelName != None:
        FlangeGeo_dict_all = tower_geo_2_dict(towerGeoExcelName)['Flange']
        flange_geo_dict = FlangeGeo_dict_all[elevation]
        mmboltType = flange_geo_dict['mmboltType']  # 螺栓规格M，用于计算螺栓疲劳等级尺寸折减系数
    else:
        mmboltType = mmboltType
    print(mmboltType)
    bolt_DC = 85
    ND = 5E6
    NA = 2E6
    gamma_f = 1.25
    if mmboltType > 30:
        ks_bolt = (30 / mmboltType) ** 0.25  # 螺栓疲劳等级尺寸折减系数
    else:
        ks_bolt = 1
    Delta_sigma_D_BOLT = ks_bolt * bolt_DC / gamma_f  * (NA / ND) ** (1 / 3)

    # 加载工作簿
    boltstressexcel_xl = openpyxl.load_workbook(boltStressExcel, data_only=True)

    # 获取所有sheet的名称
    sheet_names = boltstressexcel_xl.sheetnames
    bolt_1_sheet = boltstressexcel_xl[sheet_names[0]]
    bolt_1_sheet_max_row = bolt_1_sheet.max_row
    MKV_My = []
    for i in range(1, bolt_1_sheet_max_row, 1):
        MKV_My.append((bolt_1_sheet.cell(row=i + 1, column=1)).value)  # 单位载荷
    damage_dict = {}
    for load_name in ['Mx','My']:
        damage_dict[load_name] = {}
        for bolt_num in sheet_names:     #'''循环不同的螺栓
            damage_dict[load_name][bolt_num] = []
            for j in range(14, 22, 1):
                stress_s12 = []
                for i in range(1, bolt_1_sheet_max_row, 1):
                    stress_s12.append(boltstressexcel_xl[bolt_num].cell(row=i + 1, column=j).value)  # 单位载荷作用下应力结果
                damage = single_loadcase_damage(loc_flange_markov_dict,load_name, MKV_My,stress_s12, Delta_sigma_D_BOLT, ND)
                damage_dict[load_name][bolt_num].append(damage)
    save_bolt_fatigue_results_separate_sheets(damage_dict, outputfile)

def save_bolt_fatigue_results(results, output_file):
    """
    保存螺栓疲劳计算结果到Excel文件（按数值排序螺栓序号），写到同一个sheet里

    参数:
        results (dict): 螺栓疲劳计算结果字典，格式为 {load_name: {bolt_id: [damage_values]}}
        output_file (str): 输出文件路径
    """
    # 获取第一个载荷分量下的所有螺栓ID并排序
    first_load_name = list(results.keys())[0]
    sorted_bolt_ids = sorted(results[first_load_name].keys(),
                           key=lambda x: int(''.join(filter(str.isdigit, str(x)))))

    # 创建一个包含所有数据的DataFrame
    all_data = []
    for bolt_id in sorted_bolt_ids:
        # 为每个螺栓创建一行数据
        row_data = {'Bolt_ID': bolt_id}
        # 添加每个载荷分量在8个角度的损伤值
        for load_name in results.keys():
            for i, damage_value in enumerate(results[load_name][bolt_id]):
                row_data[f'{load_name}_Angle_{i+1}'] = damage_value
        all_data.append(row_data)

    # 创建DataFrame
    df = pd.DataFrame(all_data)

    # 保存到Excel文件
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Bolt_Fatigue_Results', index=False)

def save_bolt_fatigue_results_separate_sheets(results, output_file):
    """
    保存螺栓疲劳计算结果到Excel文件（按数值排序螺栓序号）

    参数:
        results (dict): 螺栓疲劳计算结果字典，格式为 {load_name: {bolt_id: [damage_values]}}
        output_file (str): 输出文件路径
    """
    # 获取第一个载荷分量下的所有螺栓ID并排序
    first_load_name = list(results.keys())[0]
    sorted_bolt_ids = sorted(results[first_load_name].keys(),
                           key=lambda x: int(''.join(filter(str.isdigit, str(x)))))

    # 保存到Excel文件，每个载荷分量一个sheet
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        for load_name in results.keys():
            # 创建该载荷分量的数据
            all_data = []
            for bolt_id in sorted_bolt_ids:
                # 为每个螺栓创建一行数据
                row_data = {'Bolt_ID': bolt_id}
                # 添加8个角度的损伤值
                for i, damage_value in enumerate(results[load_name][bolt_id]):
                    row_data[f'Angle_{i+1}'] = damage_value
                all_data.append(row_data)

            # 创建DataFrame并保存到对应sheet
            df = pd.DataFrame(all_data)
            df.to_excel(writer, sheet_name=f'{load_name}_Results', index=False)




if __name__ == '__main__':

    rootdir = r"D:\00.分片塔\03.分片环法兰建模\13.普通整法兰计算20251013 - 奥地利项目第2段算例-平台版本 - 螺栓后处理\1、螺栓疲劳损伤计算py程序"
    towerGeoExcelName = os.path.join(rootdir, "TowerGeoInput.xlsx")
    towerGeoExcelName = None
    mmboltType = 100
    boltstressarea = 30#螺栓应力截面积

    elevation = 0.65
    markovDir = os.path.join(rootdir, "markov")
    from bolt_stress_calculator import process_bolt_data
    process_bolt_data(os.path.join(rootdir, "Down_flage_M30_point1_my.txt"), os.path.join(rootdir, "output_analysistest.xlsx"), boltstressarea)
    #boltstressexcel = r"D:\00.分片塔\03.分片环法兰建模\13.普通整法兰计算20251013 - 奥地利项目第2段算例-平台版本 - 螺栓后处理\txt处理\output_analysis——2.xlsx"
    boltstressexcel = os.path.join(rootdir, "output_analysistest.xlsx")
    saveresultfile = os.path.join(rootdir,  "boltdamage_文件夹.xlsx")
    markovDir = os.path.join(rootdir, "markov")
    # calculate_bolt_fatigue_damage(towerGeoExcelName, markovDir, boltstressexcel, saveresultfile, elevation)
    markovDir = os.path.join(rootdir, "Markov_T_000k650.xlsx")
    saveresultfile = os.path.join(rootdir,  "boltdamage_单文件.xlsx")
    calculate_bolt_fatigue_damage(towerGeoExcelName, markovDir, boltstressexcel, saveresultfile, elevation, mmboltType)



