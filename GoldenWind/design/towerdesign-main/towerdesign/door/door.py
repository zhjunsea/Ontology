import openpyxl
import os

class DoorFem:
    def __init__(self, tower_geo_file, ult_load_file):
        self.tower_geo_file = tower_geo_file
        self.ult_load_file = ult_load_file
        self.geo_csv = []
        self.load = []
        self.order = 0
        self.initialize_csv()

    def initialize_csv(self):
        for i in range(21):
            self.geo_csv.append([0] * 8)
        for i in range(8):
            self.geo_csv[0][i] = i
        for j in range(21):
            self.geo_csv[j][0] = j

        for i in range(17):
            self.load.append([0] * 9)
        for i in range(9):
            self.load[0][i] = i
        for j in range(17):
            self.load[j][0] = j

    def remove_existing_files(self, directory):
        files_to_remove = ["design_tower_para.csv", "design_tower_para.TXT", "ultimate_load_tower_bottom.txt"]
        for file in files_to_remove:
            file_path = os.path.join(directory, file)
            if os.path.exists(file_path):
                os.remove(file_path)

    def load_excel_data(self):
        tower_gem_bk = openpyxl.load_workbook(self.tower_geo_file, data_only=True)
        ult_bk = openpyxl.load_workbook(self.ult_load_file)

        shell_gem_sht = tower_gem_bk['TowerGeo']
        flan_gem_sht = tower_gem_bk['Flange']
        door_gem_sht = tower_gem_bk['Door']
        ult_sheet = ult_bk['tower_section']

        self.shell_cell = self.get_cell_range(shell_gem_sht, 2, 2, shell_gem_sht.max_row, 6)
        self.flan_cell = self.get_cell_range(flan_gem_sht, 3, 1, flan_gem_sht.max_row, 16)
        self.door_cell = self.get_cell_range(door_gem_sht, 3, 13, 3, 22)
        self.ult_cells = self.get_cell_range(ult_sheet, 1, 1, ult_sheet.max_row, 11)

    def get_cell_range(self, sheet, start_row, start_col, end_row, end_col):
        return [[cell.value for cell in row[start_col - 1:end_col]] for row in sheet.iter_rows(min_row=start_row, max_row=end_row)]

    def find_order(self):
        for j in range(len(self.shell_cell) - 1):
            if self.shell_cell[j][0] is not None:
                if round(float(self.shell_cell[j][0]), 3) == round(float(self.flan_cell[1][0]), 3):
                    self.order = j

    def populate_geo_csv(self):
        self.geo_csv[2][1] = self.flan_cell[0][4]  # 底法兰厚度
        for row1 in range(3, self.order + 2):
            self.geo_csv[row1][1] = 1000 * (self.shell_cell[row1 - 2][0] - self.shell_cell[0][0])

        for row1 in range(1, self.order + 4):
            if self.door_cell[0][0] > self.geo_csv[row1][1]:
                if self.door_cell[0][0] < self.geo_csv[row1 + 1][1]:
                    self.geo_csv[row1 + 1][1] = self.door_cell[0][0]
                    row123 = row1 + 2

        for row1 in range(row123, self.order + 3):
            self.geo_csv[row1][1] = 1000 * (self.shell_cell[row1 - 3][0] - self.shell_cell[0][0])
        self.geo_csv[self.order + 3][1] = 1000 * (self.flan_cell[1][0] - self.flan_cell[0][0]) - self.flan_cell[1][4]  # 上法兰下表面所在标高
        self.geo_csv[self.order + 4][1] = 1000 * (self.flan_cell[1][0] - self.flan_cell[0][0])  # 上法兰上端面所在标高

        # 中径数据及壁厚数据
        for row in range(1, 4):
            self.geo_csv[row][2] = self.flan_cell[0][1] - self.flan_cell[0][5]
            self.geo_csv[row][3] = self.flan_cell[0][5]
        for row in range(4, self.order + 2):
            self.geo_csv[row][2] = self.shell_cell[row - 2][1] - self.shell_cell[row - 2][4]
            self.geo_csv[row][3] = self.shell_cell[row - 2][4]
        for row1 in range(row123, self.order + 3):
            self.geo_csv[row1][2] = self.shell_cell[row1 - 2 - 1][1] - self.shell_cell[row1 - 2 - 1][4]
            self.geo_csv[row1][3] = self.shell_cell[row1 - 2 - 1][4]
        for row in range(self.order + 3, self.order + 5):
            self.geo_csv[row][2] = self.flan_cell[1][1] - self.flan_cell[1][5]
            self.geo_csv[row][3] = self.flan_cell[1][5]

        # 塔段长度数据
        self.geo_csv[2][4] = self.geo_csv[self.order + 4][1]
        self.geo_csv[3][4] = 1
        self.geo_csv[4][4] = 1

        # 下法兰及上法兰内径
        self.geo_csv[1][5] = self.flan_cell[0][2]
        self.geo_csv[2][5] = self.flan_cell[1][2]

        # 下法兰及上法兰厚度
        self.geo_csv[1][6] = self.flan_cell[0][4]
        self.geo_csv[2][6] = self.flan_cell[1][4]

        # 塔架门洞数据
        self.geo_csv[1][7] = self.door_cell[0][0]
        self.geo_csv[5][7] = self.door_cell[0][5]
        self.geo_csv[6][7] = self.door_cell[0][6]
        self.geo_csv[7][7] = self.door_cell[0][7]
        self.geo_csv[10][7] = self.door_cell[0][4]
        self.geo_csv[11][7] = self.door_cell[0][3]
        self.geo_csv[12][7] = self.door_cell[0][8] * 0.5

        # 终止符
        for col in range(1, 4):
            self.geo_csv[self.order + 5][col] = -1

    def write_tower_para_txt(self, directory):
        tower_para = [[round(val, 3) for val in row] for row in self.geo_csv]
        file_path = os.path.join(directory, 'design_tower_para.TXT')
        with open(file_path, 'w') as f:
            for row in tower_para:
                f.write('\t'.join(map(str, row)) + '\n')

    def process_ultimate_load(self, directory):
        ult_level_list = []
        for pp in range(0, self.order):  # 底段各标高焊缝循环
            str_ult_level = self.ult_cells[18 * pp][0]
            new_str_ult_level = str_ult_level.replace('tower_section-----', "")
            ult_level_list.append(float(new_str_ult_level))

        min_level_shell = float(round(self.shell_cell[0][0], 3))

        if min_level_shell in ult_level_list:
            index = ult_level_list.index(min_level_shell)
            for i in range(1, 17):
                for j in range(1, 9):
                    self.load[i][j] = self.ult_cells[index + 1 + i][j + 1]
            file_path = os.path.join(directory, 'ultimate_load_tower_bottom.txt')
            with open(file_path, 'w') as f:
                for row in self.load:
                    f.write('\t'.join(map(str, row)) + '\n')
        else:
            print('极限载荷与塔架筒体数据的塔底标高不匹配，请检查')

    def run(self):
        input_directory = os.path.dirname(self.tower_geo_file)
        self.remove_existing_files(input_directory)
        self.load_excel_data()
        self.find_order()
        self.populate_geo_csv()
        self.write_tower_para_txt(input_directory)
        self.process_ultimate_load(input_directory)

if __name__ == '__main__':
    # 使用示例
    door_fem = DoorFem(
        tower_geo_file=r'D:\塔架计算平台\48.DNV问题\屈曲问题\door\TowerGeoInput.xlsx',
        ult_load_file=r'D:\塔架计算平台\48.DNV问题\屈曲问题\door\ultimate_loads_tower_sectionsafeFactor.xlsx'
    )
    door_fem.run()
