"""按用户口径：焊缝位置直接读 TowerGeo 的 B 列（标高(下)(m)），不推导。
逐中段列出 B 列值 → 相对段底的焊缝，并与引擎现值并排核对。"""

import openpyxl

F = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)

ws = wb["TowerGeo"]
rows = []  # (rownum, B标高下, D标高上)
for r in range(2, ws.max_row + 1):
    b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
    if b is None or d is None:
        continue
    rows.append((r, float(b), float(d)))

wf = wb["Flange"]
fl = [float(wf.cell(r, 1).value) for r in range(3, wf.max_row + 1) if wf.cell(r, 1).value is not None]
print("法兰标高:", fl)
print("中段 = 第2..5段 (法兰区间 14.89→33.37→52.97→77.89→102.53)\n")

ENG = {
    2: [215, 2800, 5600, 8400, 11200, 14000, 16680, 18305, 18480],
    3: [175, 2800, 5600, 8400, 11200, 14000, 16800, 19455, 19600],
    4: [145, 2800, 5600, 8400, 11200, 14000, 16800, 19600, 22400, 24800, 24920],
    5: [120, 2800, 5600, 8400, 11200, 14000, 16800, 19600, 22400, 24530, 24640],
}

for s in [2, 3, 4, 5]:
    lo, hi = fl[s - 1], fl[s]
    sel = [(rn, b) for (rn, b, d) in rows if lo - 1e-9 <= b < hi - 1e-9]
    base = sel[0][1]
    printed = [round((b - base) * 1000) for (rn, b) in sel]
    bcol = printed + [round((hi - base) * 1000)]   # 段顶另计
    print("=" * 74)
    print("第%d段 范围 %.3f→%.3f（段高 %d）" % (s, lo, hi, round((hi - lo) * 1000)))
    print("  B列行(r#, 标高下):", [(rn, b) for (rn, b) in sel])
    print("  B列→相对段底 (焊缝):", printed)
    print("  含段顶后          :", bcol)
    print("  引擎焊缝现值      :", ENG[s])
    print("  一致(B列+段顶)    :", bcol == ENG[s])
