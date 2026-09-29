import openpyxl
main = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(main, data_only=True)
ws = wb["TowerGeo"]
print("merged:", ws.merged_cells.ranges)
for r in [1, 2, 9, 10, 11, 18, 19]:
    print("---- row", r)
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None:
            print("   col %d (%s): %r" % (c, openpyxl.utils.get_column_letter(c), v))
