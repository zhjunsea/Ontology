"""从主体表原始 TowerGeo/Flange 重推各中段焊缝；按 W_Rung/2 口径重算附件布局。
对比两种“段底基准”：A=含贴法兰短筒节（段底=法兰标高）；B=剔除贴法兰短筒节（段底=首节顶）。"""

import openpyxl

MAIN = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(MAIN, data_only=True)

ws = wb["TowerGeo"]
courses = []  # (eb_m, et_m)
for r in range(2, ws.max_row + 1):
    eb, et = ws.cell(r, 2).value, ws.cell(r, 4).value
    if eb is None or et is None:
        continue
    courses.append((float(eb), float(et)))

wf = wb["Flange"]
flanges = [float(wf.cell(r, 1).value) for r in range(3, wf.max_row + 1) if wf.cell(r, 1).value is not None]
print("Flange 标高:", flanges, " 段数:", len(flanges) - 1)

sections = {}  # sectionNo(1-based) -> list of (eb,et)
for i in range(len(flanges) - 1):
    lo, hi = flanges[i], flanges[i + 1]
    rows = []
    for (eb, et) in courses:
        if lo - 1e-6 <= eb and et <= hi + 1e-6:
            rows.append((eb, et))
    sections[i + 1] = rows

print("\n逐段筒节（含贴法兰短节）：")
for s in sorted(sections):
    rows = sections[s]
    hs = [round((et - eb) * 1000, 1) for (eb, et) in rows]
    print("  第%d段 %s→%s  %d节 高=%s 段高=%d" %
          (s, rows[0][0], rows[-1][1], len(rows), hs, round(sum(hs))))

MID = [s for s in sorted(sections) if 2 <= s <= len(sections) - 1]
print("\n中段 =", MID)

W_RUNG = 85.0
OFF = W_RUNG / 2
FIRST, STEP = 980.0, 280.0
BAND_LO_OFF, BAND_HI_OFF = 1400.0, 1960.0


def welds_of(rows, drop_first_last):
    rs = rows[1:-1] if drop_first_last else rows
    out, acc = [], 0.0
    for (eb, et) in rs:
        acc += round((et - eb) * 1000, 1)
        out.append(acc)
    return out, acc, (rs[0][0] if rs else None)


def edge_net(h, welds):
    return min(abs(h - w) for w in welds) - OFF


def solve(platH, welds, cand):
    band_lo, band_hi = platH - BAND_HI_OFF, platH - BAND_LO_OFF
    if edge_net(FIRST, welds) <= 100:
        return None
    max_k = int((band_hi - FIRST) // STEP + 1e-9)
    reach, parent = {(1, 0)}, {}
    for c in range(2, max_k + 1):
        any_new = False
        for (cc, k) in list(reach):
            if cc != c - 1:
                continue
            for m in cand:
                k2 = k + m
                if k2 > max_k or (c, k2) in reach:
                    continue
                if edge_net(FIRST + k2 * STEP, welds) <= 100:
                    continue
                reach.add((c, k2)); parent[(c, k2)] = k; any_new = True
        if not any_new:
            break
    best, bc, bk = None, None, None
    for (c, k) in reach:
        h = FIRST + k * STEP
        if h < band_lo - 1e-6 or h > band_hi + 1e-6:
            continue
        sc = abs((h - FIRST) / (c - 1) - 1960) if c > 1 else 0
        if best is None or sc < best - 1e-9:
            best, bc, bk = sc, c, k
    if bc is None:
        return None
    hs, c, k = [], bc, bk
    while c >= 1:
        hs.insert(0, FIRST + k * STEP)
        if c == 1:
            break
        k = parent[(c, k)]; c -= 1
    return hs


for drop, tag in [(False, "基准A：含贴法兰短节（引擎现值）"), (True, "基准B：剔除贴法兰短节")]:
    print("\n" + "#" * 78)
    print(tag)
    for s in MID:
        welds, secH, base = welds_of(sections[s], drop)
        secH = round(secH)
        platH = secH - 1250
        print("  第%d段 base=%s secH=%d 焊缝=%s" % (s, base, secH, [int(round(w)) for w in welds]))
        for cand, cl in [({6, 7}, "{1680,1960}"), ({5, 6, 7}, "{1400,1680,1960}")]:
            hs = solve(platH, welds, cand)
            if hs is None:
                print("      %s → n=0 无解" % cl)
            else:
                incs = [FIRST] + [hs[i] - hs[i - 1] for i in range(1, len(hs))]
                print("      %s → n=%d 位置=%s 增量=%s 末组到平台=%.0f 最小净距=%.1f"
                      % (cl, len(hs), [int(h) for h in hs], [int(x) for x in incs],
                         platH - hs[-1], min(edge_net(h, welds) for h in hs)))
