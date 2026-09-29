import openpyxl, sys

path = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for name in ["TowerGeo", "Flange"]:
    ws = wb[name]
    print("\n===== SHEET:", name, " dims:", ws.dimensions, " max_row:", ws.max_row, " max_col:", ws.max_column)
    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            vals.append("" if v is None else str(v))
        print(f"r{r:>3}: " + " | ".join(vals))
