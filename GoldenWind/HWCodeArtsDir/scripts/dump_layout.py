import openpyxl

lay = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/项目布局表.xlsx"
main = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"

for tag, path in [("布局表", lay), ("主体表", main)]:
    wb = openpyxl.load_workbook(path, data_only=True)
    print("#" * 90)
    print(tag, "SHEETS:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print("\n===== [%s] SHEET: %s  dims:%s  max_row:%d max_col:%d" % (tag, name, ws.dimensions, ws.max_row, ws.max_column))
        for r in range(1, min(ws.max_row, 60) + 1):
            vals = []
            for c in range(1, min(ws.max_column, 14) + 1):
                v = ws.cell(r, c).value
                vals.append("" if v is None else str(v))
            line = " | ".join(vals).rstrip(" |")
            if line.strip():
                print("r%3d: %s" % (r, line))
