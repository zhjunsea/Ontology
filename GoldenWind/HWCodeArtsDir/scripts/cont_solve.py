"""若附件间距允许 [1400,1960] 内任意值（不强制 280 倍数），求各段最少附件数的排布。
目标：间距尽量大（次数最少）以节约成本；末组仍要求 ∈[1400,1960] 到平台。"""

import openpyxl

F = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)
ws = wb["TowerGeo"]
rows = []
for r in range(2, ws.max_row + 1):
    b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
    if b is None or d is None:
        continue
    rows.append((r, float(b), float(d)))
wf = wb["Flange"]
fl = [float(wf.cell(r, 1).value) for r in range(3, wf.max_row + 1) if wf.cell(r, 1).value is not None]

FIRST, EDGE, P2T = 980, 42.5, 1250
LOW, HIGH, Q = 1400, 1960, 10  # 间距下/上界, 网格步长(mm)


def welds_of(s):
    lo, hi = fl[s - 1], fl[s]
    sel = [(rn, b) for (rn, b, d) in rows if lo - 1e-9 <= b < hi - 1e-9]
    base = sel[0][1]
    return [round((b - base) * 1000) for (rn, b) in sel if abs(b - lo) > 1e-9]


def safe(h, welds):
    return all(abs(h - w) - EDGE > 100 for w in welds)


for s in [2, 3, 4, 5]:
    secH = round((fl[s] - fl[s - 1]) * 1000, 1)
    platH = secH - P2T
    welds = welds_of(s)
    bandLo, bandHi = platH - HIGH, platH - LOW
    # 末组允许落点（带内且避焊缝）
    cur = {FIRST}
    parent = {FIRST: None}
    found = None
    for c in range(2, 30):
        nxt = set()
        for p in cur:
            for d in range(LOW, HIGH + 1, Q):
                q = p + d
                if q > bandHi or q in parent or not safe(q, welds):
                    continue
                parent[q] = (p, d)
                nxt.add(q)
                if bandLo - 1e-6 <= q <= bandHi + 1e-6:
                    found = q
                    break
            if found:
                break
        if found:
            break
        cur = nxt
        if not cur:
            break
    print("=" * 84)
    print("第%d段 平台高=%.0f 平台带=[%.0f,%.0f] 焊缝=%s" % (s, platH, bandLo, bandHi, welds))
    if not found:
        print("  仍无解")
        continue
    hs, p = [], found
    while p is not None:
        hs.insert(0, p)
        p = parent[p][0] if parent[p] else None
    inc = [hs[i] - hs[i - 1] for i in range(1, len(hs))]
    print("  n=%d 位置=%s" % (len(hs), hs))
    print("        间距=%s  末组距平台=%.0f" % (inc, platH - hs[-1]))
