"""按引擎判据输出各中段附件的**具体绝对位置**与间距，并解析第3段失败点。"""

import openpyxl

F = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)
ws = wb["TowerGeo"]
rows = []
for r in range(2, ws.max_row + 1):
    b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
    if b is None or d is None:
        continue
    rows.append((r, float(b), float(d)))
wf = wb["Flange"]
fl = [float(wf.cell(r, 1).value) for r in range(3, wf.max_row + 1) if wf.cell(r, 1).value is not None]

FIRST, STEP = 980, 280.0
CANDS = [7, 6, 5]
EDGE = 42.5
PLATFORM_TO_TOP = 1250


def welds_of(s):
    lo, hi = fl[s - 1], fl[s]
    sel = [(rn, b) for (rn, b, d) in rows if lo - 1e-9 <= b < hi - 1e-9]
    base = sel[0][1]
    return [round((b - base) * 1000) for (rn, b) in sel if abs(b - lo) > 1e-9]


def safe(h, welds):
    return all(abs(h - w) - EDGE > 100 for w in welds)


def engine_dp(platH, welds, target):
    bandLo, bandHi = platH - 1960, platH - 1400
    if not safe(FIRST, welds):
        return []
    maxK = int((bandHi - FIRST) // STEP)
    if maxK < CANDS[0]:
        return []
    maxC = maxK
    reach = [[False] * (maxK + 1) for _ in range(maxC + 1)]
    par = [[-1] * (maxK + 1) for _ in range(maxC + 1)]
    reach[1][0] = True
    for c in range(2, maxC + 1):
        any_ = False
        for k in range(maxK + 1):
            if not reach[c - 1][k]:
                continue
            for m in CANDS:
                k2 = k + m
                if k2 > maxK or reach[c][k2]:
                    continue
                if not safe(FIRST + k2 * STEP, welds):
                    continue
                reach[c][k2] = True
                par[c][k2] = k
                any_ = True
        if not any_:
            break
    bestC = bestK = -1
    bestScore = 1e18
    for c in range(1, maxC + 1):
        for k in range(maxK + 1):
            if not reach[c][k]:
                continue
            h = FIRST + k * STEP
            if h < bandLo - 1e-6 or h > bandHi + 1e-6:
                continue
            score = abs(target - 1960) if c == 1 else abs((h - FIRST) / (c - 1) - target)
            if score < bestScore - 1e-9:
                bestScore, bestC, bestK = score, c, k
    if bestC < 0:
        return []
    hs, c, k = [], bestC, bestK
    while c >= 1:
        hs.insert(0, FIRST + k * STEP)
        if c == 1:
            break
        k, c = par[c][k], c - 1
    return hs


for s in [2, 3, 4, 5]:
    secH = round((fl[s] - fl[s - 1]) * 1000, 1)
    platH = secH - PLATFORM_TO_TOP
    welds = welds_of(s)
    print("=" * 90)
    print("第%d段 段高=%.0f 平台高=%.0f 平台带=[%.0f,%.0f]  焊缝=%s"
          % (s, secH, platH, platH - 1960, platH - 1400, welds))
    for tgt in (1960, 1680):
        hs = engine_dp(platH, welds, tgt)
        if not hs:
            print("  目标%d → 无解" % tgt)
            continue
        inc = [round(hs[i] - hs[i - 1]) for i in range(1, len(hs))]
        print("  目标%d → n=%d 位置=%s" % (tgt, len(hs), hs))
        print("            间距=%s  末组距平台=%.0f" % (inc, platH - hs[-1]))

# ---- 第3段失败点逐点解析 ----
s = 3
secH = round((fl[s] - fl[s - 1]) * 1000, 1)
platH = secH - PLATFORM_TO_TOP
welds = welds_of(s)
print("=" * 90)
print("第3段失败点解析：平台带 [%.0f, %.0f]，280 网格上落带内的位置：" % (platH - 1960, platH - 1400))
for k in range(50, 62):
    h = FIRST + k * STEP
    if not (platH - 1960 - 1e-6 <= h <= platH - 1400 + 1e-6):
        continue
    near = min(welds, key=lambda w: abs(w - h))
    clr = abs(near - h) - EDGE
    print("  k=%d 位置=%.0f 最近焊缝=%d 边缘净距=%.1f %s"
          % (k, h, near, clr, "OK" if clr > 100 else "★被拒(<100)"))

print()
print("对照：仅要求『相邻间距∈[1400,1960]』（不要求末组到平台∈[1400,1960]）时，")
print("从 980 贪心取大（1960→1680→1400）的位置序列：")
hs = [980]
inc = [980]
prev = 980
while True:
    for m in (1960, 1680, 1400):
        cand = prev + m
        if cand <= platH - 0 and safe(cand, welds):
            hs.append(cand)
            inc.append(m)
            prev = cand
            break
    else:
        break
    if prev >= platH - 1400:  # 已进入/越过平台带下界，停
        break
print("  位置=%s" % hs)
print("  间距=%s  (首=980)" % inc[1:])
print("  末组距平台=%.0f" % (platH - hs[-1]))
