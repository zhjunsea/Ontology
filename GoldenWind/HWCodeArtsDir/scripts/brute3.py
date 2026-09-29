"""验证：附件间距取 280 倍数、候选 {1120,1400,1680,1960} 时，第3段是否有解。
位置恒为 980+280k；末组须∈[platH-1960, platH-1400] 且避焊缝。"""

import openpyxl

F = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)
ws = wb["TowerGeo"]
rows = []
for r in range(2, ws.max_row + 1):
    b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
    if b is None or d is None:
        continue
    rows.append((r, float(b), float(d)))
wf = wb["Flange"]
fl = [float(wf.cell(r, 1).value) for r in range(3, wf.max_row + 1) if wf.cell(r, 1).value is not None]

FIRST, EDGE, P2T, STEP = 980, 42.5, 1250, 280


def welds_of(s):
    lo, hi = fl[s - 1], fl[s]
    sel = [(rn, b) for (rn, b, d) in rows if lo - 1e-9 <= b < hi - 1e-9]
    base = sel[0][1]
    return [round((b - base) * 1000) for (rn, b) in sel if abs(b - lo) > 1e-9]


def safe(h, welds):
    return all(abs(h - w) - EDGE > 100 for w in welds)


for cands in ([5, 6, 7], [4, 5, 6, 7]):
    print("候选步长(k×280) =", cands)
    for s in [2, 3, 4, 5]:
        secH = round((fl[s] - fl[s - 1]) * 1000, 1)
        platH = secH - P2T
        welds = welds_of(s)
        bandLo, bandHi = platH - 1960, platH - 1400
        reach = {(0,)}
        best = None
        # BFS over k with steps
        seen = {0: True}
        frontier = [0]
        steps_used = {0: 0}
        solutions = set()
        while frontier:
            nxt = []
            for k in frontier:
                for m in cands:
                    k2 = k + m
                    h2 = FIRST + k2 * STEP
                    if h2 > bandHi + 1e-6:
                        continue
                    if not safe(h2, welds):
                        continue
                    if k2 not in seen:
                        seen[k2] = True
                        nxt.append(k2)
                    if bandLo - 1e-6 <= h2 <= bandHi + 1e-6:
                        solutions.add(k2)
            frontier = nxt
        # also k=0 itself
        if bandLo <= FIRST <= bandHi and safe(FIRST):
            solutions.add(0)
        print("  第%d段 平台带[%.0f,%.0f] → %s"
              % (s, bandLo, bandHi, ("有解 k=" + str(sorted(solutions))) if solutions else "无解"))
    print()
