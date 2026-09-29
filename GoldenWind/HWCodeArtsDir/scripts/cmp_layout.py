"""对比老应用 ls_heights（中心距200、平台下840终止）与本引擎判据
（边缘 W_Rung/2=42.5、阈值>100、末组∈[1400,1960]）在各中段的附件排布结果。"""

import openpyxl

F = r"D:/work/Ontology/GoldenWind/design/towerdesign-main/TowerGeoInput_10563080_HH130m_6段_中径4.95m_GWH182-7.5_外置_锚栓式__主体434t-20251103-2025-11-06-11-24-50.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)
ws = wb["TowerGeo"]
rows = []
for r in range(2, ws.max_row + 1):
    b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
    if b is None or d is None:
        continue
    rows.append((r, float(b), float(d)))
wf = wb["Flange"]
fl = [float(wf.cell(r, 1).value) for r in range(3, wf.max_row + 1) if wf.cell(r, 1).value is not None]

FIRST = 980
STEP = 280.0
CANDS = [1960, 1680, 1400]
PLATFORM_TO_TOP = 1250


def welds_of(s):
    lo, hi = fl[s - 1], fl[s]
    sel = [(rn, b) for (rn, b, d) in rows if lo - 1e-9 <= b < hi - 1e-9]
    base = sel[0][1]
    return [round((b - base) * 1000) for (rn, b) in sel if abs(b - lo) > 1e-9]  # 剔除下法兰行


def old_ls(platH, welds_old, first=FIRST):
    h_abs, hv = [first], [first]
    h_prev = first
    threshold = platH - 840
    safety = 200
    while True:
        h_next = h_prev + CANDS[0]
        if h_next > threshold:
            break
        if any(abs(w - h_next) <= safety for w in welds_old):
            h_next = h_prev + CANDS[1]
            if any(abs(w - h_next) <= safety for w in welds_old):
                h_next = h_prev + CANDS[2]
        hv.append(round(h_next - h_prev))
        h_abs.append(h_next)
        h_prev = h_next
    last = sum(hv)
    if (platH - last) > 2240:
        pot = last + 1400
        if pot <= platH and not any(abs(w - pot) <= safety for w in welds_old):
            hv.append(1400)
            h_abs.append(pot)
    return h_abs, hv


def new_engine(platH, welds, edge):
    bandLo, bandHi = platH - 1960, platH - 1400

    def safe(h):
        return all(abs(h - w) - edge > 100 for w in welds)

    if not safe(FIRST):
        return []
    maxK = int((bandHi - FIRST) // STEP)
    if maxK < 7:
        return []
    reach = {(0, 1)}
    for c in range(2, maxK + 1):
        any_ = False
        for k in range(0, maxK + 1):
            if (k, c - 1) not in reach:
                continue
            for m in (7, 6, 5):
                k2 = k + m
                if k2 > maxK:
                    continue
                if (k2, c) in reach:
                    continue
                if not safe(FIRST + k2 * STEP):
                    continue
                reach.add((k2, c))
                any_ = True
        if not any_:
            break
    best = None
    for (k, c) in reach:
        h = FIRST + k * STEP
        if bandLo - 1e-6 <= h <= bandHi + 1e-6:
            if best is None or c > best[1]:
                best = (k, c)
    return [] if best is None else [FIRST + k * STEP for k in range(0, best[0] + 1)]


for s in [2, 3, 4, 5]:
    lo, hi = fl[s - 1], fl[s]
    secH = round((hi - lo) * 1000, 1)
    platH = secH - PLATFORM_TO_TOP
    welds = welds_of(s)
    welds_old = welds + [secH]
    o_abs, o_inc = old_ls(platH, welds_old)
    n_abs = new_engine(platH, welds, 42.5)
    print("=" * 78)
    print("第%d段 secH=%.1f platH=%.1f" % (s, secH, platH))
    print("  焊缝(剔法兰)=%s" % welds)
    print("  老应用 末组=%s 距平台=%.1f 组数=%d 增量=%s"
          % (o_abs[-1], platH - o_abs[-1], len(o_abs), o_inc[1:]))
    if n_abs:
        print("  本引擎 末组=%.1f 距平台=%.1f 组数=%d"
              % (n_abs[-1], platH - n_abs[-1], len(n_abs)))
    else:
        print("  本引擎 无解 n=0")
